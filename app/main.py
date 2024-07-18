from docx import Document
from odf.opendocument import load
from odf import text, teletype
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse, PlainTextResponse
from pydantic import BaseModel
from typing import Optional
import os
import io
import logging
from datetime import datetime

from app.config_loader import load_all_configs, load_config
from app.google_sheets import find_student

# Настройка логирования
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI()

courses = load_all_configs("courses")
creds_file = "credentials.json"  # Путь к файлу с учетными данными Service Account


@app.get("/courses/{course_id}/staff")
def get_course_staff(course_id: str):
    try:
        course = next((course for course in courses if course["id"] == course_id), None)
        if not course:
            raise HTTPException(status_code=404, detail="Course not found")

        config = load_config(os.path.join("courses", course["config"]))
        return config.course.staff
    except FileNotFoundError as e:
        logger.error(f"File not found: {e.detail}")
        raise HTTPException(status_code=404, detail=str(e.detail))
    except Exception as e:
        logger.error(f"Error fetching course staff: {e.detail}")
        raise HTTPException(status_code=404, detail=str(e.detail))


class Reviewer(BaseModel):
    name: str = ""
    title: str = ""


class ReportRequest(BaseModel):
    github: Optional[str] = None
    name: Optional[str] = None
    reviewer: Optional[Reviewer] = Reviewer()


@app.post("/courses/{course_id}/groups/{group_id}/labs/{lab_id}/template")
async def get_lab_template(course_id: str, group_id: str, lab_id: str, format: str, report_request: ReportRequest):
    logger.info(f"Received request: {report_request}")
    if format not in ["docx", "odt", "odf", "tex", "latex"]:
        raise HTTPException(status_code=400, detail="Invalid format")

    course = next((course for course in courses if course["id"] == course_id), None)
    if not course:
        raise HTTPException(status_code=404, detail="Course not found")

    config = load_config(os.path.join("courses", course["config"]))

    student = None
    if report_request:
        student = find_student(config.course.google['spreadsheet'], group_id, creds_file, github=report_request.github, name=report_request.name)
        if isinstance(student, dict) and "error" in student:
            return JSONResponse(status_code=400, content={"message": student["error"]})
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    template = generate_template(config, lab_id, format, student, report_request.reviewer, group_id)

    file_stream = io.BytesIO()

    if (format == 'latex' or format == 'tex'):
        format = "tex"
        file_stream.write(template.encode('utf-8'))
    else:
        template.save(file_stream)

    if format == 'odf':
        format = "odt"

    file_stream.seek(0)
    return StreamingResponse(file_stream, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document', headers={"Content-Disposition": f"attachment; filename=lab_report_{student['GitHub']}_{lab_id}.{format}"})


def generate_template(config, lab_id, format, student, reviewer, group_id):

    # Функция для преобразования полного имени в инициалы
    def get_initials(full_name):
        if (not full_name):
            return ""
        parts = full_name.split()
        initials = ' '.join([p[0] + '.' for p in parts[1:]])
        return f"{initials} {parts[0]}"

    # Определение замены меток
    replacements = {
        '*academic_position*': reviewer.title,
        '*teacher_name*': get_initials(reviewer.name),
        '*lab_number*': "№" + lab_id,
        '*course name*': config.course.alt_names[1],
        '*student group number*': group_id,
        '*initials*': get_initials(student.get('Ф.И.О.', '')),
        '*year*': str(datetime.now().year)
    }

    if format == 'docx':
        # Загрузка шаблона документа
        template_path = "templates/title_page.docx"
        document = Document(template_path)

        # Функция для замены текста в параграфах
        def replace_text_in_paragraph(paragraph, replacements):
            for key, value in replacements.items():
                if key in paragraph.text:
                    print(f"Found {key} in paragraph: {paragraph.text}")
                    paragraph.text = paragraph.text.replace(key, value)
                    print(f"Replaced {key} with {value} in paragraph")

        # Функция для замены текста в таблицах
        def replace_text_in_table(table, replacements):
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_text_in_paragraph(paragraph, replacements)

        # Обход всех параграфов в документе
        for paragraph in document.paragraphs:
            replace_text_in_paragraph(paragraph, replacements)

        # Обход всех таблиц в документе
        for table in document.tables:
            replace_text_in_table(table, replacements)

        # Добавление заголовков для лабораторной работы
        document.add_page_break()
        for heading in config.course.labs[lab_id].report:
            document.add_heading(heading + ":" + "\n", level=1)

    elif (format == 'odt' or format == 'odf'):

        template_path = "templates/title_page.odt"
        document = load(template_path)

        texts = document.getElementsByType(text.P)
        S = len(texts)

        # Замена значений в документе
        for i in range(S-len(replacements)):
            marked_text = teletype.extractText(texts[i])
            for key, value in replacements.items():
                if key in marked_text:
                    updated_text = marked_text.replace(key, value)
                    new_S = text.P()
                    new_S.setAttribute("stylename",texts[i].getAttribute("stylename"))
                    new_S.addText(updated_text)
                    texts[i].parentNode.insertBefore(new_S,texts[i])
                    texts[i].parentNode.removeChild(texts[i])
                    print(f"replaced: {marked_text} --> {value}")

        # Добавление заголовков заданий
        for heading in config.course.labs[lab_id].report:
            new_Line = text.H(outlinelevel=2)
            new_Line.addText(heading + ":")
            document.text.addElement(new_Line)

    elif (format == 'tex' or format == 'latex'):
        BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
        with open(os.path.join(BASE_DIR, 'templates', 'template.tex'), 'r', encoding='utf-8') as file:
            latex_template = file.read()

        # Замена значений в документе
        latex_template = latex_template.replace('%%REVIEWER_TITLE%%', reviewer.title)
        latex_template = latex_template.replace('%%REVIEWER_NAME%%', get_initials(reviewer.name))
        latex_template = latex_template.replace('%%LAB_NUMBER%%', lab_id)
        latex_template = latex_template.replace('%%COURSE_NAME%%', config.course.alt_names[1])
        latex_template = latex_template.replace('%%GROUP_ID%%', group_id)
        latex_template = latex_template.replace('%%STUDENT_NAME%%', get_initials(student.get('Ф.И.О.', '')))

        # Добавление заголовков заданий
        report_headers_text = '\n'.join(f'\\subsection*{{{heading}}}' for heading in config.course.labs[lab_id].report)
        latex_template = latex_template.replace('%%REPORT_HEADERS%%', report_headers_text)

        return latex_template
    else:
        # Обработка остальных форматов, если это необходимо
        raise HTTPException(status_code=400, detail="Unsupported format")

    return document


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
